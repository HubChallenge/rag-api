import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


class UnsupportedFileType(Exception):
    pass


@dataclass
class Section:
    text: str
    page: int | None = None
    heading: str | None = None


def parse_pdf(path: Path) -> list[Section]:
    reader = PdfReader(str(path))
    sections: list[Section] = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append(Section(text=text, page=i + 1))
    return sections


def parse_docx(path: Path) -> list[Section]:
    doc = DocxDocument(str(path))
    sections: list[Section] = []
    current_heading: str | None = None
    current_lines: list[str] = []

    def flush() -> None:
        text = "\n".join(current_lines).strip()
        if text:
            sections.append(Section(text=text, heading=current_heading))

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name or "") if paragraph.style else ""
        if style_name.startswith("Heading"):
            flush()
            current_heading = text
            current_lines = []
        else:
            current_lines.append(text)
    flush()

    if not sections:
        full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if full_text.strip():
            sections.append(Section(text=full_text.strip()))

    return sections


def parse_txt(path: Path) -> list[Section]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if not paragraphs:
        return []
    return [Section(text="\n\n".join(paragraphs))]


_MD_HEADING_RE = re.compile(r"^#{1,6}\s+(.*)$", re.MULTILINE)


def parse_markdown(path: Path) -> list[Section]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    matches = list(_MD_HEADING_RE.finditer(text))

    if not matches:
        stripped = text.strip()
        return [Section(text=stripped)] if stripped else []

    sections: list[Section] = []
    for i, match in enumerate(matches):
        heading = match.group(1).strip()
        start = match.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        body = text[start:end].strip()
        if body:
            sections.append(Section(text=body, heading=heading))

    if not sections:
        stripped = text.strip()
        return [Section(text=stripped)] if stripped else []

    return sections


_PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".txt": parse_txt,
    ".md": parse_markdown,
    ".markdown": parse_markdown,
}


def get_file_type(path: Path) -> str:
    ext = path.suffix.lower()
    if ext not in _PARSERS:
        raise UnsupportedFileType(f"Unsupported file type: {ext}")
    return ext.lstrip(".")


def parse_file(path: Path) -> list[Section]:
    ext = path.suffix.lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise UnsupportedFileType(f"Unsupported file type: {ext}")
    return parser(path)
